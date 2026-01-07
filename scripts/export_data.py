"""
Script pour exporter les données en Excel et PDF.
"""

import pandas as pd
from pathlib import Path
from datetime import datetime
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False


def export_to_excel(metrics_df, monthly_df, output_path=None):
    """
    Exporte les données et métriques vers un fichier Excel.
    
    Args:
        metrics_df: DataFrame avec toutes les métriques horaires
        monthly_df: DataFrame avec les agrégations mensuelles
        output_path: Chemin de sortie. Si None, utilise le chemin par défaut.
    
    Returns:
        Chemin du fichier Excel créé
    """
    if not HAS_OPENPYXL:
        raise ImportError("openpyxl n'est pas installe. Installez-le avec: pip install openpyxl")
    
    if output_path is None:
        month_str = metrics_df['datetime'].min().strftime('%Y_%m')
        output_path = Path(__file__).parent.parent / "reports" / f"export_{month_str}.xlsx"
    
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    
    # Créer un workbook
    wb = Workbook()
    
    # Style pour les en-têtes
    header_fill = PatternFill(start_color="5a9fd4", end_color="4a8bc2", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 1. Feuille : Données horaires
    ws1 = wb.active
    ws1.title = "Donnees horaires"
    
    # Préparer les données horaires
    hourly_data = metrics_df.copy()
    hourly_data['datetime'] = hourly_data['datetime'].dt.strftime('%Y-%m-%d %H:%M:%S')
    hourly_data = hourly_data[['datetime', 'price_eur_mwh', 'production_mwh', 'revenue_eur', 
                               'is_negative_price', 'negative_price_loss_eur']]
    hourly_data.columns = ['Date/Heure', 'Prix (€/MWh)', 'Production (MWh)', 
                          'Revenu (€)', 'Prix Negatif', 'Perte (€)']
    hourly_data['Prix Negatif'] = hourly_data['Prix Negatif'].map({True: 'Oui', False: 'Non'})
    
    # Écrire les en-têtes
    headers = list(hourly_data.columns)
    for col_num, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    
    # Écrire les données
    for row_num, row_data in enumerate(hourly_data.values, 2):
        for col_num, value in enumerate(row_data, 1):
            cell = ws1.cell(row=row_num, column=col_num, value=value)
            cell.border = border
            if col_num in [4, 6]:  # Colonnes monétaires
                cell.number_format = '#,##0.00'
            elif col_num == 3:  # Production
                cell.number_format = '#,##0.00'
            elif col_num == 2:  # Prix
                cell.number_format = '#,##0.00'
    
    # Ajuster la largeur des colonnes
    for col_num, header in enumerate(headers, 1):
        ws1.column_dimensions[get_column_letter(col_num)].width = max(len(str(header)), 15)
    
    # 2. Feuille : Résumé mensuel
    ws2 = wb.create_sheet("Resume mensuel")
    
    monthly_data = monthly_df.copy()
    monthly_data['year_month'] = monthly_data['year_month'].astype(str)
    monthly_data = monthly_data[['year_month', 'total_revenue_eur', 'total_negative_loss_eur',
                                 'total_production_mwh', 'avg_price_eur_mwh', 
                                 'min_price_eur_mwh', 'max_price_eur_mwh',
                                 'hours_negative_price', 'production_during_negative_mwh']]
    monthly_data.columns = ['Mois', 'Revenu Total (€)', 'Pertes Total (€)', 
                           'Production Total (MWh)', 'Prix Moyen (€/MWh)',
                           'Prix Min (€/MWh)', 'Prix Max (€/MWh)',
                           'Heures Prix Negatif', 'Production Pendant Prix Neg (MWh)']
    
    # Écrire les en-têtes
    headers2 = list(monthly_data.columns)
    for col_num, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    
    # Écrire les données
    for row_num, row_data in enumerate(monthly_data.values, 2):
        for col_num, value in enumerate(row_data, 1):
            cell = ws2.cell(row=row_num, column=col_num, value=value)
            cell.border = border
            if col_num in [2, 3]:  # Colonnes monétaires
                cell.number_format = '#,##0.00'
            elif col_num == 4:  # Production
                cell.number_format = '#,##0.00'
            elif col_num in [5, 6, 7]:  # Prix
                cell.number_format = '#,##0.00'
    
    # Ajuster la largeur des colonnes
    for col_num, header in enumerate(headers2, 1):
        ws2.column_dimensions[get_column_letter(col_num)].width = max(len(str(header)), 18)
    
    # 3. Feuille : Chiffres clés
    ws3 = wb.create_sheet("Chiffres cles")
    
    # Calculer les totaux
    total_revenue = metrics_df['revenue_eur'].sum()
    total_loss = metrics_df['negative_price_loss_eur'].sum()
    total_production = metrics_df['production_mwh'].sum()
    hours_negative = metrics_df['is_negative_price'].sum()
    total_hours = len(metrics_df)
    avg_price = metrics_df['price_eur_mwh'].mean()
    min_price = metrics_df['price_eur_mwh'].min()
    max_price = metrics_df['price_eur_mwh'].max()
    
    key_metrics = [
        ['Indicateur', 'Valeur'],
        ['Revenu Total', f'{total_revenue:,.2f} €'],
        ['Pertes Prix Negatifs', f'{total_loss:,.2f} €'],
        ['Production Totale', f'{total_production:,.2f} MWh'],
        ['Prix Moyen', f'{avg_price:,.2f} €/MWh'],
        ['Prix Minimum', f'{min_price:,.2f} €/MWh'],
        ['Prix Maximum', f'{max_price:,.2f} €/MWh'],
        ['Heures a Prix Negatif', f'{hours_negative} / {total_hours}'],
        ['Pourcentage Heures Neg', f'{hours_negative/total_hours*100:.2f}%'],
    ]
    
    for row_num, row_data in enumerate(key_metrics, 1):
        for col_num, value in enumerate(row_data, 1):
            cell = ws3.cell(row=row_num, column=col_num, value=value)
            if row_num == 1:  # En-tête
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
    
    ws3.column_dimensions['A'].width = 25
    ws3.column_dimensions['B'].width = 20
    
    # Sauvegarder
    wb.save(str(output_path))
    print(f"Export Excel genere: {output_path}")
    
    return output_path


def export_to_pdf(metrics_df, monthly_df, output_path=None):
    """
    Exporte le rapport vers un fichier PDF.
    
    Args:
        metrics_df: DataFrame avec toutes les métriques horaires
        monthly_df: DataFrame avec les agrégations mensuelles
        output_path: Chemin de sortie. Si None, utilise le chemin par défaut.
    
    Returns:
        Chemin du fichier PDF créé
    """
    if not HAS_REPORTLAB:
        raise ImportError("reportlab n'est pas installe. Installez-le avec: pip install reportlab")
    
    if output_path is None:
        month_str = metrics_df['datetime'].min().strftime('%Y_%m')
        output_path = Path(__file__).parent.parent / "reports" / f"rapport_{month_str}.pdf"
    
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    
    # Créer le document PDF
    doc = SimpleDocTemplate(str(output_path), pagesize=A4,
                          rightMargin=72, leftMargin=72,
                          topMargin=72, bottomMargin=18)
    
    # Styles
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=20,
        textColor=colors.HexColor('#2c5f7c'),
        spaceAfter=30,
        alignment=TA_CENTER
    )
    
    heading_style = ParagraphStyle(
        'CustomHeading',
        parent=styles['Heading2'],
        fontSize=14,
        textColor=colors.HexColor('#2c5f7c'),
        spaceAfter=12,
        spaceBefore=20
    )
    
    # Contenu
    story = []
    
    # Titre
    start_date = metrics_df['datetime'].min()
    end_date = metrics_df['datetime'].max()
    story.append(Paragraph("Rapport d'Analyse Énergétique", title_style))
    story.append(Spacer(1, 0.2*inch))
    story.append(Paragraph(f"Période: {start_date.strftime('%d/%m/%Y')} au {end_date.strftime('%d/%m/%Y')}", 
                          styles['Normal']))
    story.append(Paragraph(f"Date de génération: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}", 
                          styles['Normal']))
    story.append(Spacer(1, 0.3*inch))
    
    # Chiffres clés
    story.append(Paragraph("Chiffres Clés", heading_style))
    
    total_revenue = metrics_df['revenue_eur'].sum()
    total_loss = metrics_df['negative_price_loss_eur'].sum()
    total_production = metrics_df['production_mwh'].sum()
    hours_negative = metrics_df['is_negative_price'].sum()
    total_hours = len(metrics_df)
    avg_price = metrics_df['price_eur_mwh'].mean()
    min_price = metrics_df['price_eur_mwh'].min()
    max_price = metrics_df['price_eur_mwh'].max()
    
    key_data = [
        ['Indicateur', 'Valeur'],
        ['Revenu Total', f'{total_revenue:,.2f} €'],
        ['Production Totale', f'{total_production:,.2f} MWh'],
        ['Prix Moyen', f'{avg_price:,.2f} €/MWh'],
        ['Prix Minimum', f'{min_price:,.2f} €/MWh'],
        ['Prix Maximum', f'{max_price:,.2f} €/MWh'],
        ['Pertes Prix Negatifs', f'{total_loss:,.2f} €'],
        ['Heures a Prix Negatif', f'{hours_negative} / {total_hours} ({hours_negative/total_hours*100:.1f}%)'],
    ]
    
    key_table = Table(key_data, colWidths=[3*inch, 2*inch])
    key_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#5a9fd4')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 11),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafb')]),
    ]))
    story.append(key_table)
    story.append(Spacer(1, 0.3*inch))
    
    # Analyse mensuelle
    if len(monthly_df) > 0:
        story.append(Paragraph("Analyse Mensuelle", heading_style))
        
        monthly_data = []
        monthly_data.append(['Mois', 'Revenu', 'Pertes', 'Production', 'Prix Moyen', 'Heures Neg.'])
        
        for _, row in monthly_df.iterrows():
            monthly_data.append([
                str(row['year_month']),
                f'{row["total_revenue_eur"]:,.2f} €',
                f'{row["total_negative_loss_eur"]:,.2f} €',
                f'{row["total_production_mwh"]:,.2f} MWh',
                f'{row["avg_price_eur_mwh"]:,.2f} €',
                f'{int(row["hours_negative_price"])}'
            ])
        
        monthly_table = Table(monthly_data, colWidths=[1*inch, 1.2*inch, 1.2*inch, 1.2*inch, 1*inch, 0.8*inch])
        monthly_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#5a9fd4')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafb')]),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
        ]))
        story.append(monthly_table)
        story.append(Spacer(1, 0.3*inch))
    
    # Analyse des pertes
    story.append(Paragraph("Analyse des Pertes", heading_style))
    
    if total_loss > 0:
        loss_percentage = (total_loss / abs(total_revenue)) * 100 if total_revenue != 0 else 0
        story.append(Paragraph(
            f"Les pertes liées aux prix négatifs représentent <b>{total_loss:,.2f} €</b> "
            f"soit <b>{loss_percentage:.2f}%</b> du revenu total.",
            styles['Normal']
        ))
        story.append(Spacer(1, 0.1*inch))
        story.append(Paragraph(
            "Ces pertes surviennent lorsque le marché de l'électricité affiche des prix négatifs, "
            "c'est-à-dire que les producteurs doivent payer pour injecter de l'électricité sur le réseau.",
            styles['Normal']
        ))
    else:
        story.append(Paragraph(
            "Aucune perte liée aux prix négatifs n'a été enregistrée sur cette période.",
            styles['Normal']
        ))
    
    story.append(Spacer(1, 0.3*inch))
    
    # Recommandations
    story.append(Paragraph("Recommandations", heading_style))
    story.append(Paragraph("<b>Optimisation de la production:</b>", styles['Normal']))
    story.append(Paragraph(
        "• Surveillance des prix spot : Mettre en place un système d'alerte pour les périodes de prix négatifs.",
        styles['Normal']
    ))
    story.append(Paragraph(
        "• Stockage d'énergie : Envisager l'investissement dans des solutions de stockage (batteries).",
        styles['Normal']
    ))
    story.append(Paragraph(
        "• Contrats de long terme : Diversifier les revenus avec des contrats d'achat d'électricité (PPA).",
        styles['Normal']
    ))
    
    # Construire le PDF
    doc.build(story)
    print(f"Export PDF genere: {output_path}")
    
    return output_path


if __name__ == "__main__":
    # Test des exports
    from load_data import load_market_data, load_production_data
    from normalize_data import normalize_market_data, normalize_production_data
    from merge_data import merge_market_production
    from compute_metrics import compute_all_metrics
    import numpy as np
    
    print("=== Test des exports ===\n")
    
    market_df = load_market_data()
    market_norm = normalize_market_data(market_df)
    
    production_df = load_production_data()
    production_norm = normalize_production_data(production_df)
    
    merged_df = merge_market_production(market_norm, production_norm)
    
    if len(production_norm) == 0:
        print("\n[WARN] Creation de donnees de production de test...")
        merged_df['production_mwh'] = np.random.uniform(10, 50, len(merged_df))
    
    metrics_df, monthly_df = compute_all_metrics(merged_df)
    
    try:
        excel_path = export_to_excel(metrics_df, monthly_df)
        print(f"[OK] Export Excel: {excel_path}")
    except Exception as e:
        print(f"[ERREUR] Export Excel: {e}")
    
    try:
        pdf_path = export_to_pdf(metrics_df, monthly_df)
        print(f"[OK] Export PDF: {pdf_path}")
    except Exception as e:
        print(f"[ERREUR] Export PDF: {e}")

